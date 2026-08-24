#!/usr/bin/env python3
"""Convert a CSV to an XLSX with the same name, overwriting if present."""
import sys
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter

MAX_COLUMN_WIDTH = 60


def write_xlsx(csv_path: Path, xlsx_path: Path) -> None:
    df = pd.read_csv(csv_path)
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False)
        sheet = writer.sheets[writer.book.sheetnames[0]]
        for i, column in enumerate(df.columns, start=1):
            longest = max([len(str(column))] + [len(str(v)) for v in df[column]])
            sheet.column_dimensions[get_column_letter(i)].width = min(longest + 2, MAX_COLUMN_WIDTH)
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True


def main() -> None:
    if len(sys.argv) != 2:
        sys.exit("usage: csv_to_xlsx.py <path/to/file.csv>")
    csv_path = Path(sys.argv[1])
    xlsx_path = csv_path.with_suffix(".xlsx")
    write_xlsx(csv_path, xlsx_path)
    print(f"wrote {xlsx_path}")


if __name__ == "__main__":
    main()
